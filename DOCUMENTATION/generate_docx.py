import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def add_styled_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10 if level == 2 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.runs[0] if p.runs else p.add_run(text)
    if level == 1:
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy
    elif level == 2:
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 64, 175) # Blue
    elif level == 3:
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate
    return p

def convert_markdown_to_docx(md_path, docx_path, main_title="Project Documentation"):
    doc = Document()
    
    # Page setup - Normal 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows_data = []
        for l in tbl_lines:
            clean = l.strip()
            if clean.startswith('|') and clean.endswith('|'):
                cells = [c.strip() for c in clean[1:-1].split('|')]
                # Check if it's separator row
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                rows_data.append(cells)
        
        if not rows_data:
            return
        
        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for r_idx, row in enumerate(rows_data):
            for c_idx in range(num_cols):
                val = row[c_idx] if c_idx < len(row) else ""
                cell = table.cell(r_idx, c_idx)
                cell.text = val
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # Format cell text
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.name = 'Calibri'
                        r.font.size = Pt(9.5)
                        if r_idx == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            r.font.color.rgb = RGBColor(30, 41, 59)

                # Shading
                if r_idx == 0:
                    set_cell_background(cell, "1E3A8A") # Navy header
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC") # Alternating light slate
                else:
                    set_cell_background(cell, "FFFFFF")
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def flush_code_block(c_lines):
        if not c_lines:
            return
        code_text = "".join(c_lines).rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(51, 65, 85)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code_block(code_block_lines)
                code_block_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal Rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            add_styled_heading(doc, stripped[2:].strip(), 1)
        elif stripped.startswith('## '):
            add_styled_heading(doc, stripped[3:].strip(), 2)
        elif stripped.startswith('### '):
            add_styled_heading(doc, stripped[4:].strip(), 3)
        elif stripped.startswith('#### '):
            add_styled_heading(doc, stripped[5:].strip(), 4)
        # Bullet list
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            # Parse bold / inline code
            format_inline_text(p, stripped[2:])
        # Numbered list
        elif re.match(r'^\d+\.\s+', stripped):
            match = re.match(r'^\d+\.\s+', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, stripped[len(match.group(0)):])
        # Blockquote
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[2:])
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
        # Regular paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            format_inline_text(p, stripped)

        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code_block(code_block_lines)

    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")

def format_inline_text(paragraph, text):
    """Parse inline bold (**text**), italics (*text*), and code (`text`)."""
    # Simple regex token splitter
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(194, 65, 12) # Rust/Orange
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.italic = True
        else:
            paragraph.add_run(token)

if __name__ == '__main__':
    doc_dir = r"c:\Users\velan\OneDrive\Desktop\Productivity app\DOCUMENTATION"
    files = [
        ("PROJECT_REPORT.md", "PROJECT_REPORT.docx"),
        ("SYSTEM_ARCHITECTURE.md", "SYSTEM_ARCHITECTURE.docx"),
        ("DATABASE_DOCUMENTATION.md", "DATABASE_DOCUMENTATION.docx"),
        ("API_DOCUMENTATION.md", "API_DOCUMENTATION.docx"),
        ("TESTING_DOCUMENTATION.md", "TESTING_DOCUMENTATION.docx"),
        ("VIVA_NOTES.md", "VIVA_NOTES.docx"),
        ("DEPLOYMENT_GUIDE.md", "DEPLOYMENT_GUIDE.docx"),
    ]
    
    for md_file, docx_file in files:
        md_p = os.path.join(doc_dir, md_file)
        docx_p = os.path.join(doc_dir, docx_file)
        if os.path.exists(md_p):
            convert_markdown_to_docx(md_p, docx_p)
        else:
            print(f"File not found: {md_p}")

    # Generate Unified Master Document
    print("\nGenerating Unified Master Document...")
    master_md_content = []
    for md_file, _ in files:
        md_p = os.path.join(doc_dir, md_file)
        if os.path.exists(md_p):
            with open(md_p, 'r', encoding='utf-8') as f:
                master_md_content.append(f.read())
            master_md_content.append("\n\n---\n\n")

    master_md_p = os.path.join(doc_dir, "MASTER_ALL_DOCUMENTATION.md")
    master_docx_p = os.path.join(doc_dir, "PRODUCTIVITY_HUB_MASTER_DOCUMENTATION.docx")
    with open(master_md_p, 'w', encoding='utf-8') as f:
        f.write("\n".join(master_md_content))
    
    convert_markdown_to_docx(master_md_p, master_docx_p)
    if os.path.exists(master_md_p):
        os.remove(master_md_p)
    print("All Word .docx documents generated successfully!")

